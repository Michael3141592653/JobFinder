from io import BytesIO
from pathlib import Path

import docx
import pytest

from jobfinder.cv.cv_file import CvFile
from jobfinder.cv.text_extraction import (
    DATE_RANGE,
    CvUnreadableError,
    CvWithoutTextError,
    UnsupportedCvFormatError,
)

# A one-page CV made in Word and exported to PDF: a shaded sidebar (contact, skills, languages,
# education) beside the main column (profile, experience, projects, publications), with each
# job's dates right-aligned on its title line.
TWO_COLUMN_CV = Path(__file__).parent / "fixtures" / "two_column_cv.pdf"


def _pdf_with_text(text: str) -> bytes:
    """Built by hand: no PDF-writing dependency for tests."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = b"%PDF-1.4\n"
    object_offsets = []
    for object_number, object_body in enumerate(objects, start=1):
        object_offsets.append(len(pdf))
        pdf += b"%d 0 obj\n%s\nendobj\n" % (object_number, object_body)
    xref_offset = len(pdf)
    pdf += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    pdf += b"".join(b"%010d 00000 n \n" % object_offset for object_offset in object_offsets)
    pdf += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref_offset,
    )
    return pdf


def _docx_with(paragraph: str, table_cell: str) -> bytes:
    document = docx.Document()
    document.add_paragraph(paragraph)
    document.add_table(rows=1, cols=1).cell(0, 0).text = table_cell
    docx_file = BytesIO()
    document.save(docx_file)
    return docx_file.getvalue()


@pytest.mark.parametrize("file_name", ["cv.pdf", "CV.PDF"])
def test_extract_markdown_reads_a_pdf_whatever_the_extension_case(cv_text_extractor, file_name):
    markdown = cv_text_extractor.extract_markdown(
        CvFile(file_name, _pdf_with_text("Senior ML Engineer"))
    )

    assert markdown == "Senior ML Engineer"


@pytest.mark.parametrize(
    "block_in_reading_order",
    [
        pytest.param(
            "Python, SQL, Spark, Airflow, Docker, Kubernetes, FastAPI, AWS SageMaker, "
            "MLflow, CI/CD",
            id="sidebar-skills",
        ),
        pytest.param(
            "Senior Machine Learning Engineer | Mar 2021 - Present\n\nZalando, Milan (hybrid)\n\n"
            "- Led a team of 4 engineers",
            id="job-with-its-dates-and-bullets",
        ),
        pytest.param("## ML & Data", id="not-html-escaped"),
    ],
)
def test_extract_markdown_reads_a_two_column_pdf_one_column_at_a_time(
    cv_text_extractor, block_in_reading_order
):
    markdown = cv_text_extractor.extract_markdown(CvFile("cv.pdf", TWO_COLUMN_CV.read_bytes()))

    assert block_in_reading_order in markdown


@pytest.mark.parametrize(
    "title_with_its_dates",
    [
        "Senior Machine Learning Engineer | Mar 2021 - Present",
        "Machine Learning Engineer | Jun 2018 - Feb 2021",
        "Data Scientist | Sep 2016 - May 2018",
    ],
)
def test_extract_markdown_puts_right_aligned_dates_back_on_their_title(
    cv_text_extractor, title_with_its_dates
):
    markdown = cv_text_extractor.extract_markdown(CvFile("cv.pdf", TWO_COLUMN_CV.read_bytes()))

    assert title_with_its_dates in markdown


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        pytest.param("12/2025 - Present", True, id="month-number-to-present"),
        pytest.param("Mar 2021 – May 2023", True, id="month-name-en-dash"),
        pytest.param("2019 - 10/2023", True, id="year-to-month"),
        pytest.param("2020 to 2022", True, id="years-with-to"),
        pytest.param("Born 28/04/2000 in Ferrara", False, id="sentence-with-a-date"),
        pytest.param("Senior Machine Learning Engineer", False, id="title"),
        pytest.param("2021", False, id="single-year"),
    ],
)
def test_date_range_matches_only_a_whole_date_range(text, expected):
    assert (DATE_RANGE.fullmatch(text) is not None) == expected


def test_extract_markdown_reads_docx_paragraphs_and_tables(cv_text_extractor):
    file_bytes = _docx_with("Senior ML Engineer", "Python, PyTorch")

    markdown = cv_text_extractor.extract_markdown(CvFile("cv.docx", file_bytes))

    assert "Senior ML Engineer" in markdown
    assert "Python, PyTorch" in markdown


@pytest.mark.parametrize("file_name", ["cv.doc", "cv.txt", "cv"])
def test_extract_markdown_rejects_other_formats(cv_text_extractor, file_name):
    with pytest.raises(UnsupportedCvFormatError, match=r"\.pdf or \.docx"):
        cv_text_extractor.extract_markdown(CvFile(file_name, b"anything"))


def test_extract_markdown_rejects_a_damaged_file(cv_text_extractor):
    with pytest.raises(CvUnreadableError):
        cv_text_extractor.extract_markdown(CvFile("cv.pdf", b"not a pdf"))


def test_extract_markdown_rejects_a_file_without_text(cv_text_extractor):
    with pytest.raises(CvWithoutTextError):
        cv_text_extractor.extract_markdown(CvFile("blank.pdf", _pdf_with_text("")))
